
import os
import io
import uuid
import re

import pandas as pd
from docx import Document
from flask import Flask, render_template, request, redirect, url_for, flash
from werkzeug.utils import secure_filename

app = Flask(__name__)
app.secret_key = os.environ.get("SECRET_KEY", "change-this-secret-key")
app.config["MAX_CONTENT_LENGTH"] = 25 * 1024 * 1024

ALLOWED_EXTENSIONS = {".xlsx", ".xls", ".csv", ".txt", ".docx"}
MAX_FILES_IN_MEMORY = 100

# Demo index: stored in memory while the app is running.
documents = {}


def allowed_file(filename):
    return os.path.splitext(filename.lower())[1] in ALLOWED_EXTENSIONS


def extract_text(filename, raw_bytes):
    ext = os.path.splitext(filename.lower())[1]

    if ext == ".txt":
        for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
            try:
                return raw_bytes.decode(encoding)
            except UnicodeDecodeError:
                continue
        return raw_bytes.decode("utf-8", errors="replace")

    if ext == ".csv":
        try:
            frame = pd.read_csv(
                io.BytesIO(raw_bytes),
                dtype=str,
                keep_default_na=False,
                encoding="utf-8-sig"
            )
        except UnicodeDecodeError:
            frame = pd.read_csv(
                io.BytesIO(raw_bytes),
                dtype=str,
                keep_default_na=False,
                encoding="latin-1"
            )
        return frame.to_string(index=False)

    if ext in {".xlsx", ".xls"}:
        engine = "openpyxl" if ext == ".xlsx" else "xlrd"
        workbook = pd.ExcelFile(io.BytesIO(raw_bytes), engine=engine)
        parts = []

        for sheet in workbook.sheet_names:
            frame = pd.read_excel(
                workbook,
                sheet_name=sheet,
                dtype=str
            ).fillna("")

            parts.append(
                f"Sheet: {sheet}\n" + frame.to_string(index=False)
            )

        return "\n\n".join(parts)

    if ext == ".docx":
        doc = Document(io.BytesIO(raw_bytes))
        parts = [
            paragraph.text
            for paragraph in doc.paragraphs
            if paragraph.text.strip()
        ]

        for table_number, table in enumerate(doc.tables, start=1):
            parts.append(f"Table {table_number}:")

            for row in table.rows:
                parts.append(
                    " | ".join(cell.text.strip() for cell in row.cells)
                )

        return "\n".join(parts)

    raise ValueError("Unsupported file type.")


def find_matches(content, query, context=100):
    """Find every case-insensitive occurrence of a keyword or phrase."""
    if not query:
        return []

    matches = []
    pattern = re.compile(re.escape(query), re.IGNORECASE)

    for match in pattern.finditer(content):
        start, end = match.span()
        left = max(0, start - context)
        right = min(len(content), end + context)

        matches.append({
            "line": content.count("\n", 0, start) + 1,
            "position": start,
            "before": content[left:start].replace("\r", " ").replace("\n", " "),
            "matched": content[start:end],
            "after": content[end:right].replace("\r", " ").replace("\n", " "),
        })

    return matches


@app.route("/", methods=["GET", "POST"])
def index():
    if request.method == "POST":
        uploads = request.files.getlist("files")

        if not uploads or all(not f.filename for f in uploads):
            flash("Choose at least one file to upload.", "error")
            return redirect(url_for("index"))

        accepted = 0

        for uploaded in uploads:
            if not uploaded or not uploaded.filename:
                continue

            filename = secure_filename(uploaded.filename)

            if not allowed_file(filename):
                flash(f"Skipped {filename}: unsupported format.", "error")
                continue

            if len(documents) >= MAX_FILES_IN_MEMORY:
                flash("Index limit reached: maximum 100 files.", "error")
                break

            try:
                raw_bytes = uploaded.read()
                content = extract_text(filename, raw_bytes)
                doc_id = str(uuid.uuid4())

                documents[doc_id] = {
                    "id": doc_id,
                    "name": filename,
                    "ext": os.path.splitext(filename)[1].upper().lstrip("."),
                    "size": len(raw_bytes),
                    "content": content,
                    "characters": len(content),
                }

                accepted += 1

            except Exception as exc:
                flash(f"Could not read {filename}: {exc}", "error")

        if accepted:
            flash(f"Successfully indexed {accepted} file(s).", "success")

        return redirect(url_for("index"))

    query = request.args.get("q", "").strip()
    file_filter = request.args.get("type", "ALL").upper()
    sort = request.args.get("sort", "relevance")

    try:
        page = max(1, int(request.args.get("page", "1")))
    except ValueError:
        page = 1

    per_page = 25
    file_results = []
    total_matches = 0

    if query:
        for doc in documents.values():
            if file_filter != "ALL" and doc["ext"] != file_filter:
                continue

            occurrences = find_matches(doc["content"], query)

            if occurrences:
                total_matches += len(occurrences)

                file_results.append({
                    "id": doc["id"],
                    "name": doc["name"],
                    "ext": doc["ext"],
                    "characters": doc["characters"],
                    "matches": len(occurrences),
                    "occurrences": occurrences,
                })

    if sort == "name":
        file_results.sort(key=lambda item: item["name"].lower())
    else:
        file_results.sort(
            key=lambda item: (-item["matches"], item["name"].lower())
        )

    # Build a list containing every occurrence in every matching file.
    all_occurrences = []

    for doc in file_results:
        for occurrence in doc["occurrences"]:
            all_occurrences.append({
                **occurrence,
                "file_id": doc["id"],
                "name": doc["name"],
                "ext": doc["ext"],
            })

    total_occurrences = len(all_occurrences)
    total_pages = max(
        1, (total_occurrences + per_page - 1) // per_page
    )
    page = min(page, total_pages)

    start = (page - 1) * per_page
    visible_occurrences = all_occurrences[start:start + per_page]

    stats = {
        "files": len(documents),
        "characters": sum(
            doc["characters"] for doc in documents.values()
        ),
        "matching_files": len(file_results),
        "matches": total_matches,
    }

    extensions = sorted({
        doc["ext"] for doc in documents.values()
    })

    return render_template(
        "index.html",
        documents=list(documents.values()),
        file_results=file_results,
        visible_occurrences=visible_occurrences,
        query=query,
        file_filter=file_filter,
        sort=sort,
        page=page,
        total_pages=total_pages,
        total_occurrences=total_occurrences,
        stats=stats,
        extensions=extensions,
    )


@app.post("/clear")
def clear_index():
    documents.clear()
    flash("The in-memory index has been cleared.", "success")
    return redirect(url_for("index"))


@app.errorhandler(413)
def too_large(_error):
    flash("Upload too large. Keep the total request under 25 MB.", "error")
    return redirect(url_for("index"))


if __name__ == "__main__":
    port = int(os.environ.get("PORT", "5000"))
    app.run(host="0.0.0.0", port=port, debug=False)
